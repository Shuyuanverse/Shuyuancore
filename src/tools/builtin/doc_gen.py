from __future__ import annotations

import time
from pathlib import Path
from typing import Any

from src.tools.interfaces import ITool, ToolParameter, ToolResult, ToolSpec


class DocGenTool(ITool):
    """Document generation tool supporting Markdown, PDF, DOCX, and PPT output.

    Generates documents from provided content text. Markdown output is
    always supported via simple file write. PDF and DOCX generation
    require optional dependencies (reportlab and python-docx
    respectively). PPT generation is not yet implemented.
    """

    def get_spec(self) -> ToolSpec:
        """Return the tool specification.

        Returns:
            ToolSpec describing this tool's metadata and parameters.
        """

        return ToolSpec(
            name="doc_gen",
            description="Document generation tool. Supports Markdown, PDF, DOCX output.",
            category="office",
            dangerous=False,
            parameters=[
                ToolParameter(
                    name="action",
                    type="string",
                    description="Output format: markdown, pdf, docx, ppt",
                    required=True,
                ),
                ToolParameter(
                    name="content",
                    type="string",
                    description="Document content to write",
                    required=True,
                ),
                ToolParameter(
                    name="output_path",
                    type="string",
                    description="Filesystem path for the output document",
                    required=True,
                ),
                ToolParameter(
                    name="title",
                    type="string",
                    description="Optional document title",
                    required=False,
                    default="Document",
                ),
            ],
        )

    async def validate(self, params: dict[str, Any]) -> list[str]:
        """Validate parameters before execution.

        Args:
            params: Dictionary of tool parameters.

        Returns:
            List of validation error messages. Empty list means valid.
        """

        errors: list[str] = []
        action = params.get("action", "")
        valid_actions = {"markdown", "pdf", "docx", "ppt"}
        if action not in valid_actions:
            errors.append(
                f"action must be one of: {', '.join(sorted(valid_actions))}"
            )
            return errors

        content = params.get("content")
        if not content or not isinstance(content, str) or not content.strip():
            errors.append("content is required and must be a non-empty string")

        output_path = params.get("output_path")
        if not output_path or not isinstance(output_path, str) or not output_path.strip():
            errors.append("output_path is required and must be a non-empty string")

        return errors

    async def execute(
        self,
        params: dict[str, Any],
        user_id: str = "default",
    ) -> ToolResult:
        """Execute the document generation operation.

        Args:
            params: Dictionary of tool parameters.
            user_id: Identifier of the user making the request.

        Returns:
            ToolResult with the output file path on success.
        """

        action: str = params["action"]
        content: str = params["content"]
        output_path: str = params["output_path"]
        title: str = params.get("title", "Document")
        start = time.time()

        out = Path(output_path)

        try:
            if action == "markdown":
                return await self._generate_markdown(
                    content, out, start
                )
            elif action == "pdf":
                return await self._generate_pdf(
                    content, title, out, start
                )
            elif action == "docx":
                return await self._generate_docx(
                    content, title, out, start
                )
            else:
                return ToolResult(
                    success=False,
                    error="PPT generation is not yet implemented",
                    duration_ms=(time.time() - start) * 1000,
                )
        except Exception as e:
            return ToolResult(
                success=False,
                error=f"Document generation failed: {e}",
                duration_ms=(time.time() - start) * 1000,
            )

    async def _generate_markdown(
        self,
        content: str,
        output_path: Path,
        start: float,
    ) -> ToolResult:
        """Generate a Markdown document by writing content to a file.

        Args:
            content: The document content.
            output_path: Target file path.
            start: Start timestamp for duration calculation.

        Returns:
            ToolResult indicating success or failure.
        """

        try:
            output_path.parent.mkdir(parents=True, exist_ok=True)
            output_path.write_text(content, encoding="utf-8")
            return ToolResult(
                success=True,
                data={"output_path": str(output_path)},
                duration_ms=(time.time() - start) * 1000,
            )
        except OSError as e:
            return ToolResult(
                success=False,
                error=f"Failed to write Markdown file: {e}",
                duration_ms=(time.time() - start) * 1000,
            )

    async def _generate_pdf(
        self,
        content: str,
        title: str,
        output_path: Path,
        start: float,
    ) -> ToolResult:
        """Generate a PDF document using reportlab.

        Falls back with an error message if reportlab is not installed.

        Args:
            content: The document content.
            title: Document title.
            output_path: Target file path.
            start: Start timestamp for duration calculation.

        Returns:
            ToolResult indicating success or failure.
        """

        try:
            from reportlab.lib.pagesizes import A4
            from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
            from reportlab.lib.units import mm
            from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
        except ImportError:
            return ToolResult(
                success=False,
                error=(
                    "reportlab is not installed. Install it with "
                    "pip install reportlab to generate PDF documents."
                ),
                duration_ms=(time.time() - start) * 1000,
            )

        try:
            output_path.parent.mkdir(parents=True, exist_ok=True)
            doc = SimpleDocTemplate(
                str(output_path),
                pagesize=A4,
                leftMargin=20 * mm,
                rightMargin=20 * mm,
                topMargin=20 * mm,
                bottomMargin=20 * mm,
            )
            styles = getSampleStyleSheet()
            title_style = ParagraphStyle(
                "DocTitle", parent=styles["Title"], spaceAfter=12
            )
            body_style = ParagraphStyle(
                "DocBody", parent=styles["Normal"], spaceAfter=6
            )

            story: list[Any] = []
            story.append(Paragraph(title, title_style))
            story.append(Spacer(1, 6 * mm))

            for para in content.split("\n\n"):
                stripped = para.strip()
                if stripped:
                    safe_text = stripped.replace("&", "&amp;").replace(
                        "<", "&lt;"
                    ).replace(">", "&gt;").replace("\n", "<br/>")
                    story.append(Paragraph(safe_text, body_style))
                    story.append(Spacer(1, 3 * mm))

            doc.build(story)
            return ToolResult(
                success=True,
                data={"output_path": str(output_path)},
                duration_ms=(time.time() - start) * 1000,
            )
        except Exception as e:
            return ToolResult(
                success=False,
                error=f"PDF generation failed: {e}",
                duration_ms=(time.time() - start) * 1000,
            )

    async def _generate_docx(
        self,
        content: str,
        title: str,
        output_path: Path,
        start: float,
    ) -> ToolResult:
        """Generate a DOCX document using python-docx.

        Falls back with an error message if python-docx is not installed.

        Args:
            content: The document content.
            title: Document title.
            output_path: Target file path.
            start: Start timestamp for duration calculation.

        Returns:
            ToolResult indicating success or failure.
        """

        try:
            from docx import Document
            from docx.shared import Pt
        except ImportError:
            return ToolResult(
                success=False,
                error=(
                    "python-docx is not installed. Install it with "
                    "pip install python-docx to generate DOCX documents."
                ),
                duration_ms=(time.time() - start) * 1000,
            )

        try:
            output_path.parent.mkdir(parents=True, exist_ok=True)
            doc = Document()
            doc.add_heading(title, level=1)

            for para in content.split("\n\n"):
                stripped = para.strip()
                if stripped:
                    p = doc.add_paragraph(stripped)
                    for run in p.runs:
                        run.font.size = Pt(11)

            doc.save(str(output_path))
            return ToolResult(
                success=True,
                data={"output_path": str(output_path)},
                duration_ms=(time.time() - start) * 1000,
            )
        except Exception as e:
            return ToolResult(
                success=False,
                error=f"DOCX generation failed: {e}",
                duration_ms=(time.time() - start) * 1000,
            )
